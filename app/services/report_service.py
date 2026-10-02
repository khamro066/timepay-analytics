import io
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.analytics_service import get_all_employees_ranking
from app.services.notes_service import get_notes

# (data key, Excel column header, column width)
COLUMNS = [
    ("full_name", "Ism", 26),
    ("present_days", "Ishlagan kunlar soni", 14),
    ("absent_days", "Ishlamagan kunlar soni", 16),
    ("unexcused_absence_days", "Sababsiz ishga kelmagan kunlar soni", 20),
    ("excused_absence_days", "Sababli ishga kelmagan kunlar soni", 20),
    ("late_days", "Kech qolgan kunlar soni", 16),
    ("total_late_minutes", "Jami kechikkan daqiqa", 16),
    ("average_check_in_time", "O'rtacha ishga kelish vaqti", 18),
    ("early_leaving_days", "Erta ketgan kunlar soni", 16),
    ("average_check_out_time", "O'rtacha ishdan ketish vaqti", 18),
    ("total_worked_formatted", "Umumiy ish soati", 16),
    ("note", "Izoh", 32),
]


def _uzbek_worked_hours(total_minutes: int) -> str:
    return f"{total_minutes // 60} soat {total_minutes % 60} daq"


def get_report_rows(date_from: str, date_to: str, period_key: str, include_archived: bool = False) -> list[dict]:
    """Per-employee report data: the ranking fields (which already carry
    everything from _summarize_rows, computed over date_from..date_to) plus
    that employee's note for the given calendar period_key, if any. Shared
    by the on-screen Reports page and the Excel export so both always
    agree.

    include_archived surfaces paused/archived employees too — off by
    default so current reports stay focused on active staff, but available
    for a historical report that should include people no longer active."""
    ranking = get_all_employees_ranking(date_from, date_to, include_archived=include_archived)
    notes_by_employee = {n["employee_id"]: n["note"] for n in get_notes(period_key)}

    return [{**r, "note": notes_by_employee.get(r["employee_id"], "")} for r in ranking]


def build_report_workbook(
    date_from: str,
    date_to: str,
    period_key: str,
    include_archived: bool = False,
    department: str | None = None,
) -> io.BytesIO:
    rows = get_report_rows(date_from, date_to, period_key, include_archived)
    if department:
        rows = [r for r in rows if r["department"] == department]

    by_department: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        by_department[r["department"] or "Unknown"].append(r)

    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")
    dept_font = Font(bold=True, size=12)
    dept_fill = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")

    for col_idx, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row_idx = 1
    for department in sorted(by_department):
        emp_rows = by_department[department]

        ws.cell(row=row_idx, column=1, value=department)
        ws.cell(row=row_idx, column=1).font = dept_font
        for col_idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = dept_fill
        row_idx += 1

        for col_idx, (_, label, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        row_idx += 1

        for emp in sorted(emp_rows, key=lambda e: e["full_name"] or ""):
            for col_idx, (key, _, _) in enumerate(COLUMNS, start=1):
                if key == "total_worked_formatted":
                    value = _uzbek_worked_hours(emp.get("total_worked_minutes") or 0)
                elif key in ("average_check_in_time", "average_check_out_time"):
                    value = emp.get(key) or "—"
                else:
                    value = emp.get(key, "")
                ws.cell(row=row_idx, column=col_idx, value=value)
            row_idx += 1

        row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# Point widths for the same COLUMNS on a landscape A4 page (~790pt usable
# width after margins) — unrelated to the Excel sheet's character-based
# column widths above, since the two use different units entirely.
PDF_COLUMN_WIDTHS = {
    "full_name": 95,
    "present_days": 50,
    "absent_days": 52,
    "unexcused_absence_days": 58,
    "excused_absence_days": 58,
    "late_days": 50,
    "total_late_minutes": 50,
    "average_check_in_time": 56,
    "early_leaving_days": 50,
    "average_check_out_time": 58,
    "total_worked_formatted": 54,
    "note": 85,
}


def build_report_pdf(
    date_from: str,
    date_to: str,
    period_key: str,
    department: str | None = None,
    include_archived: bool = False,
) -> io.BytesIO:
    """Same report data as build_report_workbook (same columns, same
    department grouping, same optional department filter), rendered as a
    printable PDF instead of a spreadsheet."""
    rows = get_report_rows(date_from, date_to, period_key, include_archived)
    if department:
        rows = [r for r in rows if r["department"] == department]

    by_department: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        by_department[r["department"] or "Unknown"].append(r)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=28,
        title="Hisobot",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=16, alignment=0, textColor=colors.HexColor("#2E1065")
    )
    meta_style = ParagraphStyle("ReportMeta", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))
    summary_style = ParagraphStyle(
        "ReportSummary", parent=styles["Normal"], fontSize=10.5, textColor=colors.HexColor("#2E1065"), spaceBefore=6
    )
    dept_style = ParagraphStyle(
        "DeptHeader",
        parent=styles["Normal"],
        fontSize=11,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        leftIndent=2,
        spaceAfter=0,
        leading=16,
    )
    header_cell_style = ParagraphStyle("HeaderCell", parent=styles["Normal"], fontSize=7.5, textColor=colors.white, fontName="Helvetica-Bold")
    body_cell_style = ParagraphStyle("BodyCell", parent=styles["Normal"], fontSize=7.5, textColor=colors.HexColor("#222222"))

    elements: list = [
        Paragraph("Timepay Analytics — davomat hisoboti", title_style),
        Spacer(1, 4),
        Paragraph(f"Davr: {date_from} — {date_to}", meta_style),
    ]
    if department:
        elements.append(Paragraph(f"Bo'lim: {department}", meta_style))
    elements.append(Paragraph(f"Yaratilgan: {datetime.now().strftime('%Y-%m-%d %H:%M')}", meta_style))

    total_employees = len(rows)
    total_present = sum(r.get("present_days") or 0 for r in rows)
    total_expected = sum(r.get("expected_working_days") or 0 for r in rows)
    overall_rate = (total_present / total_expected) if total_expected else None
    summary_text = f"Jami xodimlar: {total_employees}"
    if overall_rate is not None:
        summary_text += f"&nbsp;&nbsp;·&nbsp;&nbsp;Umumiy davomat: {round(overall_rate * 100)}%"
    elements.append(Paragraph(summary_text, summary_style))
    elements.append(Spacer(1, 12))

    col_widths = [PDF_COLUMN_WIDTHS[key] for key, _, _ in COLUMNS]
    header_row = [Paragraph(label, header_cell_style) for _, label, _ in COLUMNS]

    for dept in sorted(by_department):
        emp_rows = sorted(by_department[dept], key=lambda e: e["full_name"] or "")

        # The department name is row 0 of the table itself (merged across
        # every column) rather than a separate Paragraph before it — a
        # Paragraph-then-Table pair can be torn apart by a page break
        # (the heading stranded at the bottom of one page, the table
        # starting fresh on the next). Making it part of the table, with
        # repeatRows=2, means the two physically cannot separate: wherever
        # the table starts, its department name travels with it, including
        # every continuation page if the department is too long for one.
        dept_name_row = [Paragraph(f"{dept} ({len(emp_rows)})", dept_style)] + [""] * (len(COLUMNS) - 1)
        table_data = [dept_name_row, header_row]
        for emp in emp_rows:
            cells = []
            for key, _, _ in COLUMNS:
                if key == "total_worked_formatted":
                    value = _uzbek_worked_hours(emp.get("total_worked_minutes") or 0)
                elif key in ("average_check_in_time", "average_check_out_time"):
                    value = emp.get(key) or "—"
                else:
                    value = emp.get(key, "")
                cells.append(Paragraph(str(value if value not in (None, "") else "—"), body_cell_style))
            table_data.append(cells)

        table = Table(table_data, colWidths=col_widths, repeatRows=2)
        style_cmds = [
            ("SPAN", (0, 0), (-1, 0)),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4C1D95")),
            ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#4C1D95")),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 1), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
            ("TOPPADDING", (0, 1), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        for i in range(2, len(table_data)):
            if i % 2 == 0:
                style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F5F3FB")))
        table.setStyle(TableStyle(style_cmds))

        elements.append(table)
        elements.append(Spacer(1, 16))

    if not rows:
        elements.append(Paragraph("Bu davr uchun ma'lumot topilmadi.", meta_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer
